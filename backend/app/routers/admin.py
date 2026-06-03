# app/routers/admin.py — full admin endpoints per PPT Phase 1, 9, 10
import io, csv
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete, func
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.database import get_db
from app.core.deps import get_current_admin
from app.core.security import hash_password
from app.models.user import User
from app.models.menu import Feedback, Menu
from app.schemas.auth import UserOut
from app.schemas.menu import FeedbackOut

router = APIRouter(prefix="/api/admin", tags=["admin"])


@router.get("/stats")
async def get_stats(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    """Dashboard summary numbers."""
    total = (await db.execute(select(func.count(User.id)).where(User.role == "trainee"))).scalar()
    voc = (await db.execute(select(func.count(User.id)).where(User.role == "trainee", User.trainee_type == "Vocational Trainee"))).scalar()
    pre = (await db.execute(select(func.count(User.id)).where(User.role == "trainee", User.trainee_type == "Pre-Trainee"))).scalar()
    menus = (await db.execute(select(func.count(Menu.id)))).scalar()
    feedbacks = (await db.execute(select(func.count(Feedback.id)))).scalar()
    avg_rating = (await db.execute(select(func.avg(Feedback.rating)))).scalar()
    now = datetime.utcnow()
    expiring_soon = (await db.execute(
        select(func.count(User.id)).where(User.role == "trainee", User.expires_at < now + timedelta(days=14), User.expires_at > now)
    )).scalar()
    return {
        "total_trainees": total, "vocational": voc, "pre_trainee": pre,
        "menus_saved": menus, "total_feedback": feedbacks,
        "avg_rating": round(float(avg_rating), 1) if avg_rating else 0,
        "expiring_soon": expiring_soon,
    }


@router.get("/users", response_model=List[UserOut])
async def list_users(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    result = await db.execute(select(User).where(User.role == "trainee").order_by(User.created_at.desc()))
    return result.scalars().all()


@router.get("/registrations")
async def list_registrations(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    """Users with their selected menus."""
    users_result = await db.execute(select(User).where(User.role == "trainee").order_by(User.created_at.desc()))
    users = users_result.scalars().all()
    out = []
    for u in users:
        menu_result = await db.execute(select(Menu).where(Menu.user_id == u.id))
        menu = menu_result.scalar_one_or_none()
        out.append({
            "id": u.id, "name": u.name, "email": u.email, "trainee_id": u.trainee_id,
            "trainee_type": u.trainee_type, "hostel_block": u.hostel_block, "mess_type": u.mess_type,
            "breakfast": menu.breakfast if menu else None,
            "lunch": menu.lunch if menu else None,
            "dinner": menu.dinner if menu else None,
            "registered_at": u.created_at.strftime("%Y-%m-%d"),
            "expires_at": u.expires_at.strftime("%Y-%m-%d") if u.expires_at else None,
        })
    return out


@router.get("/feedback", response_model=List[FeedbackOut])
async def list_feedback(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    result = await db.execute(select(Feedback).order_by(Feedback.submitted_at.desc()))
    feedbacks = result.scalars().all()
    out = []
    for fb in feedbacks:
        user_result = await db.execute(select(User).where(User.id == fb.user_id))
        user = user_result.scalar_one_or_none()
        item = FeedbackOut.model_validate(fb)
        item.user_name = user.name if user else "Unknown"
        out.append(item)
    return out


@router.post("/upload-students")
async def upload_students(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(get_current_admin),
):
    """
    PPT Phase 1: Admin bulk-uploads student database via CSV.
    CSV columns: name, email, trainee_id, trainee_type, hostel_block, mess_type, password
    """
    content = await file.read()
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(text.splitlines())
    added, skipped = 0, 0
    errors = []

    for row in reader:
        try:
            email = row["email"].strip().lower()
            tid = row["trainee_id"].strip()
            existing_email = await db.execute(select(User).where(User.email == email))
            existing_tid = await db.execute(select(User).where(User.trainee_id == tid))
            if existing_email.scalar_one_or_none() or existing_tid.scalar_one_or_none():
                skipped += 1
                continue
            trainee_type = row["trainee_type"].strip()
            expiry = datetime.utcnow() + timedelta(days=90 if trainee_type == "Vocational Trainee" else 730)
            user = User(
                name=row["name"].strip(), email=email, trainee_id=tid,
                trainee_type=trainee_type, hostel_block=row["hostel_block"].strip(),
                mess_type=row.get("mess_type", "Veg").strip(),
                password_hash=hash_password(row["password"].strip()),
                role="trainee", expires_at=expiry,
            )
            db.add(user)
            added += 1
        except Exception as e:
            errors.append(str(e))

    await db.commit()
    return {"added": added, "skipped": skipped, "errors": errors}


@router.delete("/users/expired")
async def delete_expired_users(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    """Manual trigger for expired account cleanup."""
    now = datetime.utcnow()
    result = await db.execute(select(User).where(User.role == "trainee", User.expires_at < now))
    expired = result.scalars().all()
    count = len(expired)
    for user in expired:
        await db.execute(delete(Menu).where(Menu.user_id == user.id))
        await db.execute(delete(Feedback).where(Feedback.user_id == user.id))
        await db.delete(user)
    await db.commit()
    return {"deleted": count}


@router.get("/export")
async def export_excel(db: AsyncSession = Depends(get_db), _admin: User = Depends(get_current_admin)):
    """Download all registrations as a formatted Excel file."""
    users_result = await db.execute(select(User).where(User.role == "trainee"))
    users = users_result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registrations"

    headers = ["Name","Email","Trainee ID","Type","Block","Mess Type","Breakfast","Lunch","Dinner","Registered","Expires"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D9E75")
        cell.alignment = Alignment(horizontal="center")

    for u in users:
        menu_result = await db.execute(select(Menu).where(Menu.user_id == u.id))
        menu = menu_result.scalar_one_or_none()
        ws.append([
            u.name, u.email, u.trainee_id, u.trainee_type or "-",
            u.hostel_block or "-", u.mess_type or "-",
            menu.breakfast if menu else "-",
            menu.lunch if menu else "-",
            menu.dinner if menu else "-",
            u.created_at.strftime("%Y-%m-%d"),
            u.expires_at.strftime("%Y-%m-%d") if u.expires_at else "-",
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 4, 30)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=snti_registrations.xlsx"})
